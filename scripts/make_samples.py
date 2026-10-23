"""Generates 3 sample documents (one per format) in sample_docs/. Run: python scripts/make_samples.py"""
from pathlib import Path

OUT = Path(__file__).resolve().parent.parent / "sample_docs"
OUT.mkdir(exist_ok=True)

# ------------------------------------------------------------------ 1) TXT
RAG_TXT = """Introduction to Retrieval-Augmented Generation

Retrieval-Augmented Generation, or RAG, is a technique that combines a search step with a language model. Instead of relying only on what the model memorised during training, the system first retrieves relevant passages from an external knowledge source and then asks the model to answer using those passages. RAG was popularised by a 2020 paper from Facebook AI Research.

Why use RAG?

Language models can hallucinate, meaning they produce fluent but incorrect statements. They also have a knowledge cutoff and cannot see private company documents. RAG addresses these problems by grounding answers in retrieved text. It is usually cheaper than fine-tuning because the documents can be updated without retraining the model. Another benefit is traceability: every answer can point to the source passages that supported it.

The RAG pipeline

A typical pipeline has two phases. In the indexing phase, documents are loaded, split into chunks, converted to vectors by an embedding model, and stored in a vector database. In the query phase, the user question is embedded, the nearest chunks are found by similarity search, and the top results are inserted into the prompt for the language model.

Chunking

Chunking splits long documents into small pieces. Chunks that are too large dilute the relevant facts and waste context window space. Chunks that are too small lose surrounding context. A common starting point is 500 to 1000 characters with an overlap of 10 to 20 percent, so that a sentence cut at a boundary still appears whole in the neighbouring chunk. Recursive splitting, which prefers paragraph breaks and then sentence breaks, usually keeps ideas together better than fixed-size windows.

Embeddings and vector search

An embedding model maps text to a list of numbers so that texts with similar meaning end up close together. Closeness is typically measured with cosine similarity. Popular open models include all-MiniLM-L6-v2, which produces 384-dimensional vectors, and nomic-embed-text. OpenAI offers text-embedding-3-small with 1536 dimensions. Vector databases such as ChromaDB, Pinecone and FAISS use approximate nearest neighbour indexes like HNSW to search millions of vectors quickly.

Evaluating a RAG system

Two things should be measured separately. Retrieval quality asks whether the right chunk was found; common metrics are hit rate at k and mean reciprocal rank, or MRR. Answer faithfulness asks whether the generated answer is supported by the retrieved text and does not add unsupported claims. A system can retrieve perfectly and still hallucinate, so both layers need checking.

Conversation mode

Follow-up questions such as "what about its cost?" are ambiguous on their own. A common solution is query condensation: the model rewrites the follow-up into a standalone question using the chat history, and that rewritten query is used for retrieval.
"""
(OUT / "intro_to_rag.txt").write_text(RAG_TXT, encoding="utf-8")

# ------------------------------------------------------------------ 2) DOCX
from docx import Document as Docx

POLICY = [
    ("Northwind Remote Work Policy", None),
    ("Overview", "Northwind Logistics allows eligible employees to work remotely under the rules in this policy. The policy took effect on 1 March 2025 and applies to all full-time employees who have completed their probation period."),
    ("Eligibility", "Employees must have completed six months of service to request remote work. Roles that require physical presence, such as warehouse operators and drivers, are not eligible. Managers approve requests within five business days."),
    ("Working hours and availability", "Remote employees must be reachable during core hours, which are 10:00 to 15:00 in their local time zone. Outside core hours, employees may arrange their schedule flexibly as long as they complete their contracted weekly hours of 40."),
    ("Office attendance", "Hybrid employees must come to the office at least two days per week. Team days are Tuesday and Thursday. Employees working fully remotely must attend the quarterly on-site meeting in Rotterdam."),
    ("Equipment and expenses", "Northwind provides a laptop, a monitor and a headset. Employees also receive a one-time home office stipend of 600 euros and a monthly internet allowance of 35 euros. Receipts must be submitted through the Expense Portal within 30 days of purchase."),
    ("Security requirements", "All remote work must use the company VPN and multi-factor authentication. Confidential documents must not be stored on personal devices. Lost or stolen equipment must be reported to the IT Service Desk within 24 hours."),
    ("Working from abroad", "Employees may work from another country for up to 20 days per calendar year with written manager approval. Longer stays require approval from HR because of tax and legal implications."),
    ("Contact", "Questions about this policy can be sent to people-team@northwind.example."),
]
d = Docx()
d.add_heading(POLICY[0][0], level=0)
for h, body in POLICY[1:]:
    d.add_heading(h, level=1)
    d.add_paragraph(body)
d.add_heading("Leave entitlement summary", level=1)
t = d.add_table(rows=1, cols=2)
t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Leave type", "Days per year"
for a, b in [("Annual leave", "25"), ("Sick leave", "10 paid days"), ("Parental leave", "16 weeks")]:
    r = t.add_row().cells
    r[0].text, r[1].text = a, b
d.save(OUT / "remote_work_policy.docx")

# ------------------------------------------------------------------ 3) PDF
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak

SOLAR = [
    ("The Solar System: A Short Guide", [
        "The Solar System formed about 4.6 billion years ago from a collapsing cloud of gas and dust. The Sun contains more than 99 percent of the total mass of the system. Eight planets orbit the Sun: Mercury, Venus, Earth, Mars, Jupiter, Saturn, Uranus and Neptune.",
        "The inner planets, Mercury, Venus, Earth and Mars, are rocky. The outer planets are much larger. Jupiter and Saturn are gas giants made mostly of hydrogen and helium, while Uranus and Neptune are ice giants rich in water, ammonia and methane."]),
    ("Mercury and Venus", [
        "Mercury is the smallest planet and the closest to the Sun. It completes an orbit in just 88 Earth days. It has almost no atmosphere, so temperatures swing from about 430 degrees Celsius by day to minus 180 at night.",
        "Venus is the hottest planet, with a surface temperature of roughly 465 degrees Celsius, because its thick carbon dioxide atmosphere traps heat in a runaway greenhouse effect. A day on Venus is longer than its year, and it rotates in the opposite direction to most planets."]),
    ("Earth and Mars", [
        "Earth is the only known world with liquid water on its surface and with life. Its single moon stabilises the tilt of its axis. Mars is about half the diameter of Earth and has two small moons, Phobos and Deimos.",
        "Mars shows the red colour of iron oxide dust. It hosts Olympus Mons, the tallest volcano in the Solar System at about 22 kilometres high, and Valles Marineris, a canyon system more than 4,000 kilometres long. NASA's Perseverance rover landed in Jezero Crater in February 2021 to search for signs of ancient microbial life."]),
    ("Jupiter and Saturn", [
        "Jupiter is the largest planet. Its mass is more than twice that of all the other planets combined. The Great Red Spot is a giant storm that has been observed for centuries. Jupiter has at least 90 known moons, including Ganymede, the largest moon in the Solar System.",
        "Saturn is famous for its rings, which are made of ice particles and rock. Saturn is the least dense planet, less dense than water. Its moon Titan has a thick atmosphere and lakes of liquid methane. The Cassini spacecraft orbited Saturn from 2004 until 2017."]),
    ("Uranus and Neptune", [
        "Uranus rotates on its side, with an axial tilt of about 98 degrees. It was discovered by William Herschel in 1781. Neptune was found in 1846 after mathematicians predicted its position from irregularities in the orbit of Uranus.",
        "Neptune has the strongest winds of any planet, reaching more than 2,000 kilometres per hour. Voyager 2 is the only spacecraft to have visited both Uranus and Neptune, flying past them in 1986 and 1989."]),
]
ss = getSampleStyleSheet()
story = []
for i, (h, paras) in enumerate(SOLAR):
    story.append(Paragraph(h, ss["Heading1"]))
    for p in paras:
        story += [Paragraph(p, ss["BodyText"]), Spacer(1, 10)]
    if i < len(SOLAR) - 1:
        story.append(PageBreak())
SimpleDocTemplate(str(OUT / "solar_system_guide.pdf"), pagesize=A4).build(story)
print("Wrote:", *sorted(p.name for p in OUT.iterdir()))
